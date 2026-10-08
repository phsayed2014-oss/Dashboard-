# -*- coding: utf-8 -*-
"""يبني ملف إكسل حساب مستحقات الجرد. عدّل ENTRIES يوميًا ثم شغّل: python build_jard.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

RATE_BOTH, RATE_REVIEW, RATE_COUNT = 200, 150, 100

# (التاريخ, اليوم, صيدلية 1, صيدلية 2, الموقع) — من جدول الجرد
SCHEDULE = [
    ("2026-09-01", "الثلاثاء", "15", "16", "الخرج"),
    ("2026-09-26", "السبت", "4", "5", "الخرج"),
    ("2026-09-27", "الأحد", "3", "31", "الخرج"),
    ("2026-09-28", "الاثنين", "29", "22", "الخرج"),
    ("2026-09-29", "الثلاثاء", "21", "30", "الخرج"),
    ("2026-09-30", "الأربعاء", "10", "17", "الخرج"),
    ("2026-10-01", "الخميس", "28", "", "الدلم"),
    ("2026-10-02", "الجمعة", "7", "", "الدلم"),
    ("2026-10-03", "السبت", "9", "", "الدلم"),
    ("2026-10-04", "الأحد", "32", "39", "الحوطة"),
    ("2026-10-05", "الاثنين", "33", "34", "الأفلاج"),
    ("2026-10-06", "الثلاثاء", "35", "53", "السليل"),
    ("2026-10-07", "الأربعاء", "46", "56", "تمرة - الدواسر"),
    ("2026-10-08", "الخميس", "47", "37", "الدواسر"),
    ("2026-10-09", "الجمعة", "38", "59", "الدواسر"),
    ("2026-10-10", "السبت", "36", "57", "الدواسر"),
    ("2026-10-11", "الأحد", "48", "49", "رنية"),
    ("2026-10-13", "الثلاثاء", "8", "60", "الخرج"),
    ("2026-10-14", "الأربعاء", "24", "27", "الخرج"),
    ("2026-10-15", "الخميس", "11", "40", "الخرج"),
    ("2026-10-16", "الجمعة", "18", "", "الخرج"),
    ("2026-10-17", "السبت", "19", "", "الخرج"),
    ("2026-10-18", "الأحد", "14", "", "الخرج"),
    ("2026-10-19", "الاثنين", "42", "44", "الرياض"),
    ("2026-10-20", "الثلاثاء", "43", "45", "الرياض"),
    ("2026-10-21", "الأربعاء", "23", "58", "الرياض"),
    ("2026-10-22", "الخميس", "6", "13", "الخرج"),
    ("2026-10-23", "الجمعة", "1", "", "الخرج"),
]

# السجل اليومي: (التاريخ, رقم الصيدلية, [من جرد], [من راجع])
ENTRIES = [
    ("2026-09-01", "15", ["د. حسام شلش"], ["د. حسام شلش", "د. السيد حسن", "د. أيمن حسين"]),
    ("2026-09-01", "16", ["د. حسام شلش"], ["د. حسام شلش", "د. السيد حسن", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-09-26", "4", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. علي خليفة"]),
    ("2026-09-26", "5", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. علي خليفة"]),
    ("2026-09-27", "3", ["د. علي خليفة"], ["د. علي خليفة", "د. حسام شلش", "د. السيد حسن"]),
    ("2026-09-27", "31", ["د. علي خليفة"], ["د. علي خليفة", "د. حسام شلش", "د. السيد حسن", "د. أيمن حسين"]),
    ("2026-09-28", "29", ["د. حسام شلش"], ["د. حسام شلش", "د. السيد حسن", "د. علي خليفة", "د. أيمن حسين"]),
    ("2026-09-28", "22", ["د. حسام شلش"], ["د. حسام شلش", "د. السيد حسن", "د. علي خليفة"]),
    ("2026-09-29", "21", ["د. علي خليفة"], ["د. السيد حسن", "د. علي خليفة", "د. حسام شلش", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-09-29", "30", ["د. علي خليفة"], ["د. السيد حسن", "د. علي خليفة", "د. حسام شلش", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-09-30", "10", ["د. حسام شلش"], ["د. حسام شلش", "د. علي خليفة", "د. السيد حسن", "د. أيمن حسين"]),
    ("2026-09-30", "17", ["د. حسام شلش"], ["د. حسام شلش", "د. علي خليفة", "د. السيد حسن", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-10-01", "28", ["د. علي خليفة"], ["د. حسام شلش", "د. السيد حسن", "د. علي خليفة", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-10-02", "7", ["د. حسام شلش"], ["د. حسام شلش", "د. السيد حسن", "د. علي خليفة", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-10-03", "9", ["د. علي خليفة"], ["د. حسام شلش", "د. السيد حسن", "د. علي خليفة", "د. أيمن حسين", "د. الحسن الغنيمي"]),
    ("2026-10-04", "32", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-04", "39", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-05", "33", ["د. السيد حسن"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-05", "34", ["د. السيد حسن"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-06", "35", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-06", "53", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-07", "46", ["د. السيد حسن"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-07", "56", ["د. السيد حسن"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-08", "47", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
    ("2026-10-08", "37", ["د. حسام شلش"], ["د. السيد حسن", "د. حسام شلش", "د. الحسن الغنيمي", "د. محمد جمال"]),
]

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FONT = Font(name=FONT, color="0000FF")
BASE = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0 "ر.س";-#,##0;"-"'


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CENTER, BORDER
        ws.column_dimensions[c.column_letter].width = w


def cellfmt(c, font=BASE, fmt=None):
    c.font, c.alignment, c.border = font, CENTER, BORDER
    if fmt:
        c.number_format = fmt


wb = Workbook()

# ---------- الإعدادات ----------
st = wb.active
st.title = "الإعدادات"
st.sheet_view.rightToLeft = True
header(st, 1, ["البند", "المبلغ (ريال)"], [30, 16])
for r, (k, v) in enumerate([("جرد + مراجعة", RATE_BOTH), ("مراجعة فقط", RATE_REVIEW), ("جرد فقط", RATE_COUNT)], 2):
    cellfmt(st.cell(row=r, column=1, value=k))
    c = st.cell(row=r, column=2, value=v)
    cellfmt(c, INPUT_FONT, MONEY)
    c.fill = PatternFill("solid", fgColor="FFFF00")
st["B2"].comment = Comment("المبالغ حسب تعليمات المستخدم. المبلغ محسوب لكل صيدلية لكل شخص.", "Claude")
st["A6"] = "ملاحظات:"
st["A6"].font = BOLD
notes = [
    "المستحق يُحسب لكل صيدلية: الشخص اللي جرد وراجع نفس الصيدلية = 200، راجع فقط = 150، جرد فقط = 100.",
    "الخلايا الصفراء (باللون الأزرق) قابلة للتعديل، وكل الحسابات تتحدث تلقائيًا.",
    "في شيت (السجل) كل صف = شخص واحد في صيدلية واحدة، واكتب (نعم) في عمود جرد و/أو راجع.",
]
for i, n in enumerate(notes, 7):
    st.cell(row=i, column=1, value="• " + n).font = BASE
st.column_dimensions["A"].width = 30

# ---------- جدول الجرد ----------
sc = wb.create_sheet("جدول الجرد")
sc.sheet_view.rightToLeft = True
header(sc, 1, ["اليوم", "التاريخ", "الصيدلية 1", "الصيدلية 2", "الموقع"], [12, 14, 12, 12, 18])
for r, (d, day, p1, p2, loc) in enumerate(SCHEDULE, 2):
    vals = [day, d, f"صيدلية رقم {p1}", f"صيدلية رقم {p2}" if p2 else "***", loc]
    for col, v in enumerate(vals, 1):
        cellfmt(sc.cell(row=r, column=col, value=v))
sc.freeze_panes = "A2"
SCHED_LAST = len(SCHEDULE) + 1

# ---------- السجل ----------
lg = wb.create_sheet("السجل")
lg.sheet_view.rightToLeft = True
header(lg, 1, ["التاريخ", "اليوم", "الموقع", "رقم الصيدلية", "الاسم", "جرد", "راجع", "نوع المهمة", "المستحق"],
       [13, 11, 16, 12, 20, 8, 8, 16, 13])
rows = []
for d, ph, counters, reviewers in ENTRIES:
    for name in dict.fromkeys(counters + reviewers):
        rows.append((d, ph, name, "نعم" if name in counters else "", "نعم" if name in reviewers else ""))
LOG_MAX = 1000
for r in range(2, LOG_MAX + 1):
    data = rows[r - 2] if r - 2 < len(rows) else None
    if data:
        d, ph, name, cnt, rev = data
        for col, v in zip([1, 4, 5, 6, 7], [d, int(ph), name, cnt, rev]):
            cellfmt(lg.cell(row=r, column=col, value=v), INPUT_FONT)
    lg.cell(row=r, column=2, value=f'=IF(A{r}="","",IFERROR(INDEX(\'جدول الجرد\'!$A$2:$A${SCHED_LAST},MATCH(A{r},\'جدول الجرد\'!$B$2:$B${SCHED_LAST},0)),""))')
    lg.cell(row=r, column=3, value=f'=IF(A{r}="","",IFERROR(INDEX(\'جدول الجرد\'!$E$2:$E${SCHED_LAST},MATCH(A{r},\'جدول الجرد\'!$B$2:$B${SCHED_LAST},0)),""))')
    lg.cell(row=r, column=8, value=f'=IF(E{r}="","",IF(AND(F{r}="نعم",G{r}="نعم"),"جرد + مراجعة",IF(F{r}="نعم","جرد فقط",IF(G{r}="نعم","مراجعة فقط",""))))')
    lg.cell(row=r, column=9, value=f'=IF(H{r}="جرد + مراجعة",الإعدادات!$B$2,IF(H{r}="مراجعة فقط",الإعدادات!$B$3,IF(H{r}="جرد فقط",الإعدادات!$B$4,0)))')
    for col in range(1, 10):
        c = lg.cell(row=r, column=col)
        if col in (2, 3, 8, 9):
            cellfmt(c, BASE, MONEY if col == 9 else None)
        elif not data:
            cellfmt(c, INPUT_FONT)
dv = DataValidation(type="list", formula1='"نعم"', allow_blank=True)
lg.add_data_validation(dv)
dv.add(f"F2:G{LOG_MAX}")
lg.freeze_panes = "A2"
lg.auto_filter.ref = f"A1:I{LOG_MAX}"

# ---------- الملخص ----------
sm = wb.create_sheet("الملخص", 0)
sm.sheet_view.rightToLeft = True
sm["A1"] = "ملخص مستحقات الجرد"
sm["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E78")
header(sm, 3, ["الاسم", "جرد + مراجعة", "مراجعة فقط", "جرد فقط", "عدد الصيدليات", "إجمالي المستحق"], [22, 14, 14, 12, 14, 16])
names = list(dict.fromkeys(n for *_, n, _, _ in rows))
NAME_SLOTS = max(15, len(names) + 5)
L = f"السجل!$H$2:$H${LOG_MAX}"
N = f"السجل!$E$2:$E${LOG_MAX}"
for i in range(NAME_SLOTS):
    r = 4 + i
    cellfmt(sm.cell(row=r, column=1, value=names[i] if i < len(names) else None), INPUT_FONT)
    for col, kind in zip([2, 3, 4], ["جرد + مراجعة", "مراجعة فقط", "جرد فقط"]):
        cellfmt(sm.cell(row=r, column=col, value=f'=IF(A{r}="","",COUNTIFS({N},A{r},{L},"{kind}"))'))
    cellfmt(sm.cell(row=r, column=5, value=f'=IF(A{r}="","",SUM(B{r}:D{r}))'))
    cellfmt(sm.cell(row=r, column=6, value=f'=IF(A{r}="","",SUMIFS(السجل!$I$2:$I${LOG_MAX},{N},A{r}))'), BOLD, MONEY)
tr = 4 + NAME_SLOTS
cellfmt(sm.cell(row=tr, column=1, value="الإجمالي"), BOLD)
for col in range(2, 7):
    L_ = sm.cell(row=4, column=col).column_letter
    cellfmt(sm.cell(row=tr, column=col, value=f"=SUM({L_}4:{L_}{tr-1})"), BOLD, MONEY if col == 6 else None)
    sm.cell(row=tr, column=col).fill = PatternFill("solid", fgColor="DDEBF7")
sm.cell(row=tr, column=1).fill = PatternFill("solid", fgColor="DDEBF7")
cellfmt(sm.cell(row=tr + 2, column=1, value="للتحقق: إجمالي السجل"), BASE)
cellfmt(sm.cell(row=tr + 2, column=6, value=f"=SUM(السجل!$I$2:$I${LOG_MAX})"), BASE, MONEY)
sm.cell(row=2, column=1, value="اكتب أي اسم جديد في عمود الاسم (أزرق) وسيُحسب تلقائيًا من شيت السجل").font = Font(name=FONT, italic=True, color="7F7F7F")

# ---------- حسب اليوم ----------
dy = wb.create_sheet("حسب اليوم", 1)
dy.sheet_view.rightToLeft = True
dates = sorted(dict.fromkeys(d for d, *_ in ENTRIES))
day_of = {d: day for d, day, *_ in SCHEDULE}
header(dy, 1, ["الاسم"] + [f"{day_of.get(d, '')}\n{d}" for d in dates] + ["الإجمالي"],
       [22] + [13] * len(dates) + [15])
dy.row_dimensions[1].height = 32
D = f"السجل!$A$2:$A${LOG_MAX}"
for i, name in enumerate(names):
    r = 2 + i
    cellfmt(dy.cell(row=r, column=1, value=f"=الملخص!A{4 + i}"), BOLD)
    for j, d in enumerate(dates):
        col = 2 + j
        dy.cell(row=1, column=col).value = f"{day_of.get(d, '')}\n{d}"
        dy.cell(row=1000, column=col, value=d)  # مفتاح التاريخ (مخفي)
        cl = dy.cell(row=1000, column=col).column_letter
        cellfmt(dy.cell(row=r, column=col, value=f'=SUMIFS(السجل!$I$2:$I${LOG_MAX},{N},$A{r},{D},{cl}$1000)'), BASE, MONEY)
    last = dy.cell(row=r, column=1 + len(dates)).column_letter
    cellfmt(dy.cell(row=r, column=2 + len(dates), value=f"=SUM(B{r}:{last}{r})"), BOLD, MONEY)
tr2 = 2 + len(names)
cellfmt(dy.cell(row=tr2, column=1, value="الإجمالي"), BOLD)
for col in range(2, 3 + len(dates)):
    L_ = dy.cell(row=2, column=col).column_letter
    c = dy.cell(row=tr2, column=col, value=f"=SUM({L_}2:{L_}{tr2 - 1})")
    cellfmt(c, BOLD, MONEY)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
dy.cell(row=tr2, column=1).fill = PatternFill("solid", fgColor="DDEBF7")
dy.row_dimensions[1000].hidden = True
dy.freeze_panes = "B2"

OUT = "inventory/jard_payroll.xlsx"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", len(rows), "rows")

# حفظ النتائج المحسوبة داخل الملف حتى تظهر في برامج العرض على الموبايل
try:
    import os, sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import cache_values
    cache_values.embed(OUT)
    print("cached values embedded")
except ImportError as e:
    print("skip caching:", e)

# ---------- ملف السجل فقط ----------
LOG_OUT = "inventory/jard_log.xlsx"
lw = Workbook()
ls = lw.active
ls.title = "السجل"
ls.sheet_view.rightToLeft = True
header(ls, 1, ["التاريخ", "اليوم", "الموقع", "رقم الصيدلية", "الاسم", "جرد", "راجع", "نوع المهمة", "المستحق"],
       [13, 11, 16, 12, 20, 8, 8, 16, 13])
loc_of = {d: loc for d, _, _, _, loc in SCHEDULE}
for r, (d, ph, name, cnt, rev) in enumerate(rows, 2):
    kind = "جرد + مراجعة" if cnt and rev else ("جرد فقط" if cnt else "مراجعة فقط")
    amt = {"جرد + مراجعة": RATE_BOTH, "مراجعة فقط": RATE_REVIEW, "جرد فقط": RATE_COUNT}[kind]
    for col, v in enumerate([d, day_of.get(d, ""), loc_of.get(d, ""), int(ph), name, cnt, rev, kind, amt], 1):
        cellfmt(ls.cell(row=r, column=col, value=v), BASE, MONEY if col == 9 else None)
tr3 = len(rows) + 2
for col in range(1, 10):
    c = ls.cell(row=tr3, column=col)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    cellfmt(c, BOLD)
ls.cell(row=tr3, column=1, value="الإجمالي")
ls.cell(row=tr3, column=9, value=f"=SUM(I2:I{tr3 - 1})").number_format = MONEY
ls.freeze_panes = "A2"
ls.auto_filter.ref = f"A1:I{tr3 - 1}"
lw.calculation.fullCalcOnLoad = True
lw.save(LOG_OUT)
try:
    cache_values.embed(LOG_OUT)
except NameError:
    pass
print("saved log-only file")
