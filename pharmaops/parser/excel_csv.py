"""Excel / CSV ingestion for CLI pipeline."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .normalize import normalize_row


def parse_excel_path(path: str | Path) -> list[dict[str, str]]:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, str]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        raw = {headers[i]: r[i] for i in range(len(headers)) if i < len(r)}
        norm = normalize_row(raw)
        if norm:
            rows.append(norm)
    wb.close()
    return rows


def parse_csv_path(path: str | Path) -> list[dict[str, str]]:
    import csv

    path = Path(path)
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows: list[dict[str, str]] = []
        for raw in reader:
            norm = normalize_row(raw)
            if norm:
                rows.append(norm)
        return rows
